from __future__ import annotations

import io
import unittest
from unittest.mock import patch
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from manager_api.main import TEMPLATE_SHEETS, analyse_rows, app, normalized_header

BASE = Path(__file__).resolve().parents[1]


class Phase37ManagerApplicationTests(unittest.TestCase):
    def test_manager_is_not_a_placeholder_shell(self):
        html = (BASE / "manager_app" / "index.html").read_text(encoding="utf-8")
        js = (BASE / "manager_app" / "app.js").read_text(encoding="utf-8")
        api = (BASE / "manager_api" / "main.py").read_text(encoding="utf-8")
        for label in (
            "Nueva temporada", "Nuevo equipo", "Nuevo partido", "Añadir usuario existente",
            "Descargar plantilla Excel", "Nueva versión",
        ):
            self.assertIn(label, html)
        for endpoint in (
            '/api/seasons', '/api/competitions', '/api/teams', '/api/players',
            '/api/matches', '/api/members', '/api/themes', '/api/import/analyze',
        ):
            self.assertIn(endpoint, api)
        self.assertIn("submitModal", js)
        self.assertIn("setLoginBusy", js)
        self.assertIn("forgot-password", html)

    def test_health_and_index_are_current_and_never_cached(self):
        with TestClient(app) as client:
            health = client.get("/api/health")
            self.assertEqual(health.status_code, 200)
            self.assertEqual(health.json()["build"], "phase38")
            self.assertIn("no-store", health.headers["cache-control"])
            index = client.get("/")
            self.assertEqual(index.status_code, 200)
            self.assertIn("SecretariatPro Manager", index.text)


    def test_login_endpoint_returns_a_real_connected_manager_session(self):
        import manager_api.main as manager

        class FakeClient:
            def __init__(self):
                self.user_id = "user-1"
            def login(self, email, password):
                if password != "correcta":
                    raise RuntimeError("Credenciales incorrectas")
            def workspace_context(self):
                return {
                    "workspaces": [{"id": "workspace-1", "name": "Asociación", "workspace_type": "association", "role": "owner"}],
                    "workspace": {"id": "workspace-1", "name": "Asociación", "workspace_type": "association", "role": "owner"},
                    "subscription": {
                        "id": "sub-1", "status": "active", "plan_id": "association_test",
                        "starts_at": "2026-01-01T00:00:00+00:00", "ends_at": "2027-12-31T00:00:00+00:00",
                        "offline_grace_days": 4,
                        "plan": {"id": "association_test", "name": "Asociación de pruebas", "entitlements": {}},
                    },
                }
            def user_profile(self):
                return {"id": "user-1", "email": "a.garciarenones@example.com", "display_name": "Alejandro"}
            def logout(self):
                pass

        manager.runtime.logout()
        with patch.object(manager, "ScoreboardSupabaseClient", FakeClient), TestClient(manager.app) as client:
            response = client.post("/api/auth/login", json={"email": "a.garciarenones@example.com", "password": "correcta"})
            self.assertEqual(response.status_code, 200)
            payload = response.json()
            self.assertTrue(payload["connected"])
            self.assertTrue(payload["permissions"]["write_allowed"])
            self.assertEqual(payload["workspace"]["workspace_type"], "association")
        manager.runtime.logout()

    def test_template_contains_every_bulk_import_scope(self):
        with TestClient(app) as client:
            response = client.get("/api/import/template.xlsx")
            self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content), read_only=True)
        self.assertEqual(set(workbook.sheetnames), set(TEMPLATE_SHEETS))
        for title, headers in TEMPLATE_SHEETS.items():
            row = next(workbook[title].iter_rows(values_only=True))
            self.assertEqual(["is_coach" if value == "Entrenador" else value for value in row], headers)

    def test_column_mapping_understands_common_spanish_headers(self):
        self.assertEqual(normalized_header("Dorsal"), "shirt_number")
        self.assertEqual(normalized_header("N.º camiseta"), "shirt_number")
        result = analyse_rows("Jugadores", [["Nombre", "Apellidos", "Dorsal"], ["Ana", "García", 17]])
        self.assertEqual(result["row_count"], 1)
        self.assertEqual(result["preview"][0]["shirt_number"], 17)

    def test_manager_sql_can_add_existing_members_safely(self):
        sql = (BASE / "SUPABASE_PHASE37_MANAGER.sql").read_text(encoding="utf-8")
        self.assertIn("sp_add_existing_member", sql)
        self.assertIn("sp_has_workspace_role", sql)
        self.assertIn("auth.users", sql)
        self.assertIn("scoreboard_operators", sql)

    def test_manager_client_has_real_crud_methods(self):
        source = (BASE / "supabase_client.py").read_text(encoding="utf-8")
        for method in (
            "manager_create_season", "manager_create_competition", "manager_create_team",
            "manager_create_player", "manager_create_match", "manager_add_existing_member",
            "manager_create_theme", "manager_publish_theme",
        ):
            self.assertIn(f"def {method}", source)


if __name__ == "__main__":
    unittest.main()
