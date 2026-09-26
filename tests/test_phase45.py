from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from manager_api.main import (
    TEMPLATE_SHEETS,
    build_player_import_template,
    download_import_template,
)

BASE = Path(__file__).resolve().parents[1]


class Phase45Tests(unittest.IsolatedAsyncioTestCase):
    def test_release_and_assets_are_phase45(self):
        import manager_release
        self.assertGreaterEqual(int(manager_release.MANAGER_RELEASE.replace("phase", "")), 45)
        html = (BASE / "manager_app" / "index.html").read_text(encoding="utf-8")
        self.assertIn(f"styles.css?v={manager_release.MANAGER_RELEASE}", html)
        self.assertIn(f"app.js?v={manager_release.MANAGER_RELEASE}", html)

    def test_player_template_is_a_valid_single_sheet_workbook(self):
        from io import BytesIO
        workbook = load_workbook(BytesIO(build_player_import_template()))
        self.assertEqual(workbook.sheetnames, ["Jugadores"])
        headers = [cell.value for cell in workbook["Jugadores"][1]]
        self.assertEqual(["is_coach" if value == "Entrenador" else value for value in headers], TEMPLATE_SHEETS["Jugadores"])

    async def test_desktop_download_writes_the_excel_file(self):
        with tempfile.TemporaryDirectory() as directory:
            with patch("manager_api.main.manager_downloads_directory", return_value=Path(directory)):
                result = await download_import_template()
            destination = Path(result["path"])
            self.assertTrue(destination.exists())
            self.assertGreater(destination.stat().st_size, 1000)
            self.assertEqual(destination.suffix, ".xlsx")

    def test_settings_and_languages_exist(self):
        html = (BASE / "manager_app" / "index.html").read_text(encoding="utf-8")
        js = (BASE / "manager_app" / "app.js").read_text(encoding="utf-8")
        css = (BASE / "manager_app" / "styles.css").read_text(encoding="utf-8")
        self.assertIn('data-view="settings"', html)
        self.assertIn('id="settings-language"', html)
        self.assertIn('id="settings-theme"', html)
        for language in ("es", "en", "sv", "cs", "fi", "de"):
            self.assertIn(f'value="{language}"', html)
        self.assertIn("downloadPlayerTemplate", js)
        self.assertIn('/api/import/scoped/${$("#import-kind").value}/template/download', js)
        self.assertIn('html[data-app-theme="light"]', css)
        self.assertIn('.sidebar-footer .logout', css)


if __name__ == "__main__":
    unittest.main()
